import base64
from io import BytesIO

from odoo import fields, models, _
from odoo.exceptions import UserError
from openpyxl import load_workbook


class ImportSaleOrderUpdate(models.TransientModel):
    _name = 'import.sale.order.update'
    _description = 'Import Fix Order Amount'

    file = fields.Binary(string="File Excel (xlsx)", required=True)
    file_name = fields.Char()

    def import_file_ver(self):
        if not self.file:
            raise UserError(_("Vui lòng chọn file Excel"))

        if not self.file_name.endswith(('.xlsx', '.xls')):
            raise UserError(_("Chỉ hỗ trợ file Excel"))

        excel_data = BytesIO(base64.b64decode(self.file))
        workbook = load_workbook(filename=excel_data, data_only=True)
        sheet = workbook.active


        errors = []
        updated = 0

        for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            order_name = str(row[0]).strip() if row[0] else False
            total_excel = row[1]
            total_tax_excel = row[2]

            if not order_name or not total_excel:
                continue

            total_excel = float(total_excel)

            order = self.env['sale.order'].search([('name', '=', order_name)], limit=1)
            # POS:
            # order = self.env['pos.order'].search([('pos_reference', '=', order_name)], limit=1)
            if not order:
                errors.append(f"Dòng {i}: Không tìm thấy đơn {order_name}")
                continue
            # đánh dấu để không đồng bộ lại
            # if order.mark_done_cron:
            #     continue
            order.write({
                "mark_done_cron" : True,
            })
            lines = order.order_line
            if not lines:
                continue

            total_current = sum(lines.mapped('price_subtotal'))

            if total_current == 0:
                continue
            # if abs(total_current - total_excel) < 1:
            #     continue
            total_diff =  abs(total_current - total_excel)
            order.write({
                "total_diff" : total_diff,
            })
            # xem tỉ lệ chênh lệch
            ratio = total_excel / total_current

            new_subtotals = []
            for line in lines:
                new_price = line.price_unit * ratio
                line.write({'price_unit': new_price})
                new_subtotals.append(line.product_uom_qty * new_price)
            # Nếu tổng tiền mới vẫn bị lệch thì lấy phần dư cộng vào đơn giá của sản phẩm cuối cùng
            diff = total_excel - sum(new_subtotals)
            if abs(diff) > 0:
                last_line = lines[-1]
                qty = last_line.product_uom_qty or 1
                last_line.write({
                    'price_unit': last_line.price_unit + (diff / qty)
                })

            total_tax_order = 0
            for line in lines:
                total_tax_order += round(line.price_tax, 0)
            if abs(total_tax_order - total_tax_excel) != 0:
                diff_tax = total_tax_order - total_tax_excel
                order.write({
                    "is_used_custom_tax": True
                })
                to_update_line = lines.filtered(lambda line: line.price_tax and line.price_tax > diff_tax)
                if to_update_line:
                    to_update_line[0].write({
                        'price_tax': to_update_line[0].price_tax - diff_tax
                    })

            updated += 1

        message = f"Đã cập nhật {updated} đơn"
        if errors:
            message += "\nLỗi:\n" + "\n".join(errors[:10])

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Kết quả'),
                'message': message,
                'type': 'success' if not errors else 'warning',
            }
        }
